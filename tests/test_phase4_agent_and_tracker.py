import os
import tempfile
from pathlib import Path
import pytest
import docx

from jobpilot.agent.graph import run_job_agent
from jobpilot.db.database import DatabaseManager
from jobpilot.db.repository import JobPilotRepository
from jobpilot.models.schemas import ApplicationStatus, CandidateProfile, Job
from jobpilot.resume.parser import ResumeParser
from jobpilot.sources.demo_data import DemoJobSource
from jobpilot.tracker.service import ApplicationTrackerService


def test_resume_parser_docx():
    parser = ResumeParser()
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tf:
        docx_path = Path(tf.name)

    # Generate a small sample docx
    doc = docx.Document()
    doc.add_heading("Rahul Girase", 0)
    doc.add_paragraph("AI & Machine Learning student graduating in 2027.")
    doc.add_paragraph("Skills: Python, Scikit-learn, LangGraph, FastAPI, SQL.")
    doc.add_paragraph("Projects: MHT-CET College Predictor, Air Drawing System using OpenCV and MediaPipe.")
    doc.save(str(docx_path))

    try:
        parsed = parser.parse_file(docx_path)
        assert "Python" in parsed["parsed_skills"]
        assert "LangGraph" in parsed["parsed_skills"]
        assert "FastAPI" in parsed["parsed_skills"]
        assert len(parsed["parsed_projects"]) >= 1
    finally:
        if docx_path.exists():
            os.remove(docx_path)


def test_tracker_approval_and_duplicate_prevention(monkeypatch):
    with tempfile.NamedTemporaryFile(suffix=".db", delete=False) as tf:
        temp_path = Path(tf.name)

    manager = DatabaseManager(temp_path)
    repo = JobPilotRepository(manager)
    # Monkeypatch singleton repository in tracker service
    monkeypatch.setattr("jobpilot.tracker.service.repository", repo)

    job = Job(
        id="job_app_test_1",
        title="ML Engineer",
        company="InnoTech",
        location="Pune",
        description="ML Engineer entry level.",
        url="https://example.com/innotech",
        source_name="Demo",
    )
    repo.upsert_job(job)

    tracker = ApplicationTrackerService()

    # 1. Approve
    app = tracker.approve_application("job_app_test_1", "Looks great!")
    assert app.status == ApplicationStatus.APPROVED

    # 2. Record Submission
    app_sub = tracker.record_manual_submission("job_app_test_1", channel="LinkedIn", confirmation_id="CONF-888")
    assert app_sub.status == ApplicationStatus.SUBMITTED

    # 3. Duplicate submission attempt must be prevented
    app_dup = tracker.record_manual_submission("job_app_test_1", channel="Company Site")
    assert app_dup.status == ApplicationStatus.SUBMITTED
    assert "CONF-888" in app_dup.submission_notes  # Original submission details preserved

    # 4. Reject test
    tracker.reject_application("job_app_test_1", reason="No longer interested")
    app_rej = repo.get_application("job_app_test_1")
    assert app_rej.status == ApplicationStatus.REJECTED

    if temp_path.exists():
        try:
            os.remove(temp_path)
        except Exception:
            pass


def test_langgraph_pipeline_run():
    # Run agent with curated demo jobs
    demo_jobs = DemoJobSource().fetch_jobs(limit=4)
    result = run_job_agent(custom_jobs=demo_jobs)

    assert result["status"] == "completed"
    assert "summary" in result
    summary = result["summary"]
    assert summary["total_jobs_discovered"] >= 3
    assert len(result["ranked_matches"]) >= 3
    # High match jobs should have drafts created
    assert len(result["drafts"]) >= 1
